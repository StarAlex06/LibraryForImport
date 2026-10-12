from collections.abc import Iterator

import pandas as pd
from openpyxl import load_workbook

from table_importer.base import DataSource
from table_importer.exceptions import (
    DataImportError,
    ScanError,
    SheetNotFoundError,
    SourceNotFoundError,
)
from table_importer.metadata import ColumnMetadata, SheetMetadata


class ExcelSource(DataSource):
    """Data source implementation for Excel XLSX files."""

    def _check_file(self) -> None:
        if not self.path.exists():
            raise SourceNotFoundError(
                f"File does not exist: {self.path}"
            )

        if not self.path.is_file():
            raise SourceNotFoundError(
                f"Path is not a file: {self.path}"
            )

    def list_sheets(self) -> list[str]:
        """Return names of all worksheets in the workbook."""
        self._check_file()

        try:
            with pd.ExcelFile(self.path, engine="openpyxl") as workbook:
                return workbook.sheet_names
        except Exception as exc:
            raise ScanError(
                f"Could not list Excel sheets in {self.path}: {exc}"
            ) from exc

    def _check_sheet_name(self, sheet_name: str) -> None:
        if sheet_name not in self.list_sheets():
            raise SheetNotFoundError(
                f"Excel source has no sheet named {sheet_name!r}."
            )

    def scan(
        self,
        sheet_name: str,
        sample_size: int = 1000,
    ) -> SheetMetadata:
        """Inspect up to sample_size rows of an Excel worksheet."""
        self._check_file()

        if sample_size < 1:
            raise ValueError("sample_size must be greater than zero.")

        self._check_sheet_name(sheet_name)

        try:
            data = pd.read_excel(
                self.path,
                sheet_name=sheet_name,
                engine="openpyxl",
                nrows=sample_size + 1,
            )
        except Exception as exc:
            raise ScanError(
                f"Could not scan Excel sheet {sheet_name!r}: {exc}"
            ) from exc

        scanned_data = data.iloc[:sample_size]
        scanned_rows = len(scanned_data)
        all_rows_scanned = len(data) <= sample_size

        columns = tuple(
            ColumnMetadata(
                name=str(column),
                data_type=str(scanned_data[column].dtype),
                missing_count=int(scanned_data[column].isna().sum()),
                scanned_count=scanned_rows,
            )
            for column in scanned_data.columns
        )

        return SheetMetadata(
            name=sheet_name,
            row_count=scanned_rows if all_rows_scanned else None,
            column_count=len(scanned_data.columns),
            columns=columns,
            scanned_rows=scanned_rows,
            row_count_is_exact=all_rows_scanned,
            missing_counts_are_exact=all_rows_scanned,
        )

    def read_chunk(
        self,
        sheet_name: str,
        offset: int,
        chunk_size: int,
    ) -> pd.DataFrame:
        """Read a worksheet chunk starting at the given data-row offset."""
        self._check_file()

        if offset < 0:
            raise ValueError("offset cannot be negative.")

        if chunk_size < 1:
            raise ValueError("chunk_size must be greater than zero.")

        self._check_sheet_name(sheet_name)

        try:
            return pd.read_excel(
                self.path,
                sheet_name=sheet_name,
                engine="openpyxl",
                skiprows=range(1, offset + 1),
                nrows=chunk_size,
            )
        except Exception as exc:
            raise DataImportError(
                f"Could not read Excel sheet {sheet_name!r}: {exc}"
            ) from exc
    def iter_chunks(
        self,
        sheet_name: str,
        chunk_size: int,
        offset: int = 0,
    ) -> Iterator[pd.DataFrame]:
        """Stream worksheet rows in chunks using openpyxl read-only mode."""
        self._check_file()

        if offset < 0:
            raise ValueError("offset cannot be negative.")

        if chunk_size < 1:
            raise ValueError("chunk_size must be greater than zero.")

        self._check_sheet_name(sheet_name)

        workbook = None

        try:
            # Read the header separately. pandas handles empty worksheets
            # more consistently than openpyxl's read-only row iterator.
            empty_frame = pd.read_excel(
                self.path,
                sheet_name=sheet_name,
                engine="openpyxl",
                nrows=0,
            )
            headers = list(empty_frame.columns)

            workbook = load_workbook(
                self.path,
                read_only=True,
                data_only=True,
            )
            worksheet = workbook[sheet_name]
            rows = worksheet.iter_rows(values_only=True)

            # Skip the header row, if the worksheet has one.
            next(rows, None)

            # Skip already imported data rows.
            for _ in range(offset):
                if next(rows, None) is None:
                    yield empty_frame
                    return


            batch = []
            pending_blank_rows = []

            for row in rows:
                if all(value is None for value in row):
                    # Temporarily keep blank rows. They may be inside
                    # the data, so we cannot discard them immediately.
                    pending_blank_rows.append(row)
                    continue

                # A non-empty row follows: blank rows before it are
                # internal data rows and must be preserved.
                for blank_row in pending_blank_rows:
                    batch.append(blank_row)

                    if len(batch) == chunk_size:
                        yield pd.DataFrame(batch, columns=headers)
                        batch = []

                pending_blank_rows.clear()

                batch.append(row)

                if len(batch) == chunk_size:
                    yield pd.DataFrame(batch, columns=headers)
                    batch = []

            # Emit any remaining data rows. Blank rows at the very end
            # are intentionally ignored because they may be Excel padding.
            if batch:
                yield pd.DataFrame(batch, columns=headers)
            elif offset == 0:
                # An empty worksheet should still preserve its columns.
                yield empty_frame


        except Exception as exc:
            raise DataImportError(
                f"Could not stream Excel sheet {sheet_name!r}: {exc}"
            ) from exc

        finally:
            if workbook is not None:
                workbook.close()